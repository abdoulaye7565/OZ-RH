"""Tests de l'assistant documentaire (prompt 6.2, chapitre 16.2.1/16.3.2 du CDC).

Aucun appel réel aux services externes (Anthropic, Voyage) : `assistance_api_cle`
et `assistance_voyage_api_cle` valent None par défaut dans tous les
environnements de test (jamais de secret réel dans les tests, CLAUDE.md point
10) — les appels externes sont simulés, la logique de découpage, d'extraction
et de récupération est, elle, entièrement réelle."""
import io

import docx
import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.core.config import settings
from app.core.security import creer_access_token
from app.models.document import Document
from app.models.enums import ConfidentialiteDocument, StatutDocument
from app.models.segment_document import SegmentDocument
from app.services.assistance import assistant_documentaire, decoupage, embeddings, extraction, indexation
from app.services.assistance.embeddings import ResultatEmbeddings
from app.services.document_service import approuver, creer_document, nouvelle_version, soumettre_approbation
from app.schemas.document import DocumentCreation


def _entete(utilisateur):
    jeton = creer_access_token(utilisateur.id, utilisateur.role.value)
    return {"Authorization": f"Bearer {jeton}"}


# --- Découpage ---


def test_decouper_regroupe_les_paragraphes_sous_la_taille_cible():
    texte = "Un paragraphe.\nUn deuxième paragraphe."
    segments = decoupage.decouper(texte, taille_cible=1000)
    assert segments == ["Un paragraphe.\nUn deuxième paragraphe."]


def test_decouper_scinde_au_dela_de_la_taille_cible():
    texte = "a" * 10 + "\n" + "b" * 10
    segments = decoupage.decouper(texte, taille_cible=15)
    assert len(segments) == 2


def test_decouper_texte_vide():
    assert decoupage.decouper("") == []


# --- Extraction ---


def test_extraire_docx_regroupe_par_titre(tmp_path):
    document = docx.Document()
    document.add_heading("Consignes de sécurité", level=1)
    document.add_paragraph("Toujours porter le harnais.")
    document.add_heading("Procédure d'urgence", level=1)
    document.add_paragraph("Appeler le référent SHEQ.")
    chemin = tmp_path / "test.docx"
    document.save(chemin)

    sections = extraction.extraire_sections(str(chemin))

    assert dict(sections) == {
        "Consignes de sécurité": "Toujours porter le harnais.",
        "Procédure d'urgence": "Appeler le référent SHEQ.",
    }


def test_extraire_xlsx_une_section_par_feuille(tmp_path):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Registre"
    feuille.append(["Danger", "Criticité"])
    feuille.append(["Chute", "12"])
    chemin = tmp_path / "test.xlsx"
    classeur.save(chemin)

    sections = extraction.extraire_sections(str(chemin))

    assert len(sections) == 1
    assert sections[0][0] == "Registre"
    assert "Chute" in sections[0][1]


def test_extraire_fichier_inexistant_ne_leve_pas():
    assert extraction.extraire_sections("chemin/inexistant.docx") == []


def test_extraire_extension_non_reconnue():
    assert extraction.extraire_sections("fichier.txt") == []


# --- Embeddings (service externe simulé) ---


def test_embeddings_sans_cle_ne_leve_pas_et_journalise(db_session):
    resultat = embeddings.calculer(db_session, fonction="test", textes=["x"], type_entree="document")
    assert resultat.disponible is False


def test_embeddings_refuse_un_contenu_bytes(db_session):
    from app.services.assistance.garde_fous import ContenuRefuseError

    with pytest.raises(ContenuRefuseError):
        embeddings.calculer(db_session, fonction="test", textes=[b"secret"], type_entree="document")


# --- Indexation ---


def test_indexer_document_sans_fichier_ne_produit_aucun_segment(db_session, referent_sheq):
    document = Document(
        reference="POL-SHEQ-999", intitule="Test", niveau=1, version="01",
        redacteur_id=referent_sheq.id, statut=StatutDocument.EN_VIGUEUR, fichier=None,
    )
    db_session.add(document)
    db_session.commit()

    assert indexation.indexer_document(db_session, document) is False
    assert db_session.scalars(select(SegmentDocument)).first() is None


def test_indexer_document_avec_embeddings_indisponibles_ne_produit_aucun_segment(db_session, referent_sheq, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path))
    fichier = docx.Document()
    fichier.add_heading("Section", level=1)
    fichier.add_paragraph("Un texte quelconque.")
    (tmp_path / "documents").mkdir()
    fichier.save(tmp_path / "documents" / "test.docx")

    document = Document(
        reference="POL-SHEQ-998", intitule="Test", niveau=1, version="01",
        redacteur_id=referent_sheq.id, statut=StatutDocument.EN_VIGUEUR, fichier="documents/test.docx",
    )
    db_session.add(document)
    db_session.commit()

    # assistance_voyage_api_cle reste None (défaut de test) : le service
    # d'embeddings est indisponible, comme en conditions réelles sans clé.
    assert indexation.indexer_document(db_session, document) is False
    assert db_session.scalars(select(SegmentDocument)).first() is None


def test_indexer_document_avec_embeddings_simules_cree_des_segments(db_session, referent_sheq, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path))
    fichier = docx.Document()
    fichier.add_heading("Consignes de sécurité", level=1)
    fichier.add_paragraph("Toujours porter le harnais avant de monter.")
    (tmp_path / "documents").mkdir()
    fichier.save(tmp_path / "documents" / "test.docx")

    document = Document(
        reference="POL-SHEQ-997", intitule="Test", niveau=1, version="01",
        redacteur_id=referent_sheq.id, statut=StatutDocument.EN_VIGUEUR, fichier="documents/test.docx",
    )
    db_session.add(document)
    db_session.commit()

    monkeypatch.setattr(
        embeddings, "calculer",
        lambda db, *, fonction, textes, type_entree, utilisateur_id=None: ResultatEmbeddings(
            disponible=True, vecteurs=[[1.0, 0.0] for _ in textes]
        ),
    )

    assert indexation.indexer_document(db_session, document) is True
    segments = db_session.scalars(select(SegmentDocument).where(SegmentDocument.document_id == document.id)).all()
    assert len(segments) == 1
    assert segments[0].section == "Consignes de sécurité"


def test_retirer_segments_supprime_bien_les_lignes(db_session):
    segment = SegmentDocument(document_id=1, section="X", texte="x", embedding=[1.0], position=0)
    db_session.add(segment)
    db_session.commit()

    indexation.retirer_segments(db_session, 1)

    assert db_session.scalars(select(SegmentDocument)).first() is None


# --- Approbation d'un document : déclenche/retire l'indexation ---


def test_approbation_indexe_et_nouvelle_version_retire_les_segments_de_l_ancienne(
    db_session, referent_sheq, responsable, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path))
    monkeypatch.setattr(
        embeddings, "calculer",
        lambda db, *, fonction, textes, type_entree, utilisateur_id=None: ResultatEmbeddings(
            disponible=True, vecteurs=[[1.0, 0.0] for _ in textes]
        ),
    )

    def _fichier_docx(texte: str) -> str:
        fichier = docx.Document()
        fichier.add_heading("Section", level=1)
        fichier.add_paragraph(texte)
        (tmp_path / "documents").mkdir(exist_ok=True)
        import uuid
        nom = f"{uuid.uuid4().hex}.docx"
        fichier.save(tmp_path / "documents" / nom)
        return f"documents/{nom}"

    doc_v1 = creer_document(
        db_session, DocumentCreation(reference="POL-SHEQ-996", intitule="Test", niveau=1),
        _fichier_docx("Version 1 du texte."), redacteur_id=referent_sheq.id,
    )
    soumettre_approbation(db_session, doc_v1, modifie_par_id=referent_sheq.id)
    approuver(db_session, doc_v1, approbateur_id=responsable.id)

    segments_v1 = db_session.scalars(select(SegmentDocument).where(SegmentDocument.document_id == doc_v1.id)).all()
    assert len(segments_v1) == 1

    doc_v2 = nouvelle_version(db_session, doc_v1, _fichier_docx("Version 2 du texte."), redacteur_id=referent_sheq.id)
    soumettre_approbation(db_session, doc_v2, modifie_par_id=referent_sheq.id)
    approuver(db_session, doc_v2, approbateur_id=responsable.id)

    # La version 1, désormais archivée, n'a plus de segments.
    assert db_session.scalars(select(SegmentDocument).where(SegmentDocument.document_id == doc_v1.id)).first() is None
    segments_v2 = db_session.scalars(select(SegmentDocument).where(SegmentDocument.document_id == doc_v2.id)).all()
    assert len(segments_v2) == 1


# --- Assistant documentaire (bout en bout, appels externes simulés) ---


def test_repondre_fonction_desactivee_par_defaut(db_session, technicien):
    resultat = assistant_documentaire.repondre(db_session, question="Une question ?", utilisateur_id=technicien.id)
    assert resultat.disponible is False


def test_repondre_sans_segment_pertinent_ne_genere_pas_de_reponse(db_session, technicien, monkeypatch):
    monkeypatch.setattr(settings, "assistance_activee", True)
    monkeypatch.setattr(settings, "assistance_assistant_documentaire_active", True)
    monkeypatch.setattr(
        embeddings, "calculer",
        lambda db, *, fonction, textes, type_entree, utilisateur_id=None: ResultatEmbeddings(
            disponible=True, vecteurs=[[1.0, 0.0]]
        ),
    )
    appele = {"valeur": False}

    def _appeler_espion(*args, **kwargs):
        appele["valeur"] = True

    from app.services.assistance import client as assistance_client
    monkeypatch.setattr(assistance_client, "appeler", _appeler_espion)

    resultat = assistant_documentaire.repondre(
        db_session, question="Quel est le budget prévisionnel 2027 ?", utilisateur_id=technicien.id
    )

    assert resultat.disponible is True
    assert resultat.reponse_trouvee is False
    assert resultat.reponse == assistant_documentaire.MESSAGE_ABSENCE_DE_REPONSE
    assert appele["valeur"] is False  # le modèle de génération n'est jamais appelé


def test_repondre_avec_segment_pertinent_cite_le_document(db_session, technicien, referent_sheq, monkeypatch):
    monkeypatch.setattr(settings, "assistance_activee", True)
    monkeypatch.setattr(settings, "assistance_assistant_documentaire_active", True)

    document = Document(
        reference="FOR-SHEQ-004", intitule="Fiche SLAM", niveau=4, version="01",
        redacteur_id=referent_sheq.id, statut=StatutDocument.EN_VIGUEUR, fichier="documents/x.docx",
    )
    db_session.add(document)
    db_session.commit()
    segment = SegmentDocument(
        document_id=document.id, section="Vent et conditions météo",
        texte="En cas de vent supérieur à 50 km/h, l'intervention est reportée.",
        embedding=[1.0, 0.0], position=0,
    )
    db_session.add(segment)
    db_session.commit()

    monkeypatch.setattr(
        embeddings, "calculer",
        lambda db, *, fonction, textes, type_entree, utilisateur_id=None: ResultatEmbeddings(
            disponible=True, vecteurs=[[1.0, 0.0]],
        ),
    )
    from app.services.assistance import client as assistance_client
    from app.services.assistance.client import ResultatAppel
    monkeypatch.setattr(
        assistance_client, "appeler",
        lambda db, *, fonction, prompt, utilisateur_id: ResultatAppel(
            disponible=True,
            contenu="Il faut reporter l'intervention (FOR-SHEQ-004, Vent et conditions météo).",
        ),
    )

    resultat = assistant_documentaire.repondre(
        db_session, question="Que faire si le vent se lève pendant une intervention en hauteur ?",
        utilisateur_id=technicien.id,
    )

    assert resultat.disponible is True
    assert resultat.reponse_trouvee is True
    assert len(resultat.references) == 1
    assert resultat.references[0].reference == "FOR-SHEQ-004"
    assert resultat.references[0].section == "Vent et conditions météo"


# --- Route ---


def test_route_question_documentaire_fonction_desactivee(client, technicien):
    reponse = client.post(
        "/api/v1/assistance/question-documentaire",
        headers=_entete(technicien),
        json={"question": "Une question ?"},
    )
    assert reponse.status_code == 200
    assert reponse.json()["disponible"] is False


def test_route_fichier_document_absent_pour_document_sans_fichier(client, referent_sheq):
    doc = client.post(
        "/api/v1/documents", headers=_entete(referent_sheq),
        data={"reference": "POL-SHEQ-995", "intitule": "Test", "niveau": "1"},
    ).json()
    reponse = client.get(f"/api/v1/documents/{doc['id']}/fichier", headers=_entete(referent_sheq))
    assert reponse.status_code == 404
