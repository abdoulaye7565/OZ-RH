"""Extraction du texte des documents en vigueur (prompt 6.2, chapitre 16.3.2
du CDC). Un document sans fichier joint, ou d'un format non reconnu, ne lève
aucune exception : il est simplement exclu de l'indexation (cohérent avec "le
mode dégradé est la norme" — l'approbation d'un document ne doit jamais
échouer à cause d'un problème d'extraction)."""
from pathlib import Path

import docx
import pypdf
from openpyxl import load_workbook


def extraire_sections(chemin: str) -> list[tuple[str, str]]:
    """Renvoie une liste de (étiquette de section, texte). Étiquette : titre de
    section pour un .docx (regroupé sous le dernier titre "Heading" rencontré),
    "Page N" pour un PDF, nom de feuille pour un .xlsx."""
    extension = Path(chemin).suffix.lower()
    try:
        if extension == ".docx":
            return _extraire_docx(chemin)
        if extension == ".xlsx":
            return _extraire_xlsx(chemin)
        if extension == ".pdf":
            return _extraire_pdf(chemin)
    except Exception:
        # Fichier corrompu, protégé par mot de passe, etc. : traité comme un
        # document non indexable, pas comme une erreur de l'application —
        # voir la même approche pour l'absence de fichier.
        return []
    return []


def _extraire_docx(chemin: str) -> list[tuple[str, str]]:
    document = docx.Document(chemin)
    sections: list[tuple[str, str]] = []
    section_courante = "Introduction"
    paragraphes_courants: list[str] = []

    def _cloturer():
        texte = "\n".join(paragraphes_courants).strip()
        if texte:
            sections.append((section_courante, texte))

    for paragraphe in document.paragraphs:
        if paragraphe.style.name.startswith("Heading") and paragraphe.text.strip():
            _cloturer()
            section_courante = paragraphe.text.strip()
            paragraphes_courants = []
        elif paragraphe.text.strip():
            paragraphes_courants.append(paragraphe.text.strip())
    _cloturer()
    return sections


def _extraire_xlsx(chemin: str) -> list[tuple[str, str]]:
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    sections: list[tuple[str, str]] = []
    for feuille in classeur.worksheets:
        lignes = []
        for ligne in feuille.iter_rows(values_only=True):
            valeurs = [str(v) for v in ligne if v is not None]
            if valeurs:
                lignes.append(" | ".join(valeurs))
        texte = "\n".join(lignes).strip()
        if texte:
            sections.append((feuille.title, texte))
    return sections


def _extraire_pdf(chemin: str) -> list[tuple[str, str]]:
    lecteur = pypdf.PdfReader(chemin)
    sections: list[tuple[str, str]] = []
    for numero, page in enumerate(lecteur.pages, start=1):
        texte = (page.extract_text() or "").strip()
        if texte:
            sections.append((f"Page {numero}", texte))
    return sections
