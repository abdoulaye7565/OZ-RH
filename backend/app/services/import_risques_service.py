"""Import initial du registre des risques (prompt 4.1) depuis un gabarit CSV/XLSX,
sur le modèle du fichier réel docs/REG-SHEQ-001_Registre_des_Risques.xlsx (feuille
"Registre", en-têtes en ligne 4).

Colonnes du gabarit réel sans équivalent dans le modèle RISQUE/COTATION_RISQUE :
« Responsable », « Échéance », « Statut » — ce sont en réalité des champs
d'ACTION (section 5.2.5 : "Actions avec... responsable, échéance... statut"),
pas de RISQUE. Le module Actions existe déjà (prompt 1.2, `Action.risque_id`)
mais « Responsable » dans le classeur est un intitulé de rôle ou de service
("Référent SHEQ", "Direction", "Tout le personnel"), pas un identifiant
d'utilisateur réel : aucune correspondance fiable n'est possible sans deviner.
Ces trois colonnes sont donc acceptées si présentes mais ignorées, comme pour
les colonnes non modélisées du gabarit d'inventaire du parc (prompt 3.1) — une
action de suivi peut être créée manuellement ensuite via l'API Actions.

« C » et « Niveau de risque » sont également ignorées à l'import : ce sont des
valeurs calculées (règle 6, CLAUDE.md), jamais importées telles quelles, toujours
recalculées côté serveur à partir de P et G.
"""
import csv
import io
from datetime import date

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.models.enums import CategorieRisque
from app.schemas.risque import (
    CotationEntree,
    LigneErreurImportRisque,
    RapportImportRisques,
    RisqueCreation,
    RisqueSortie,
)
from app.services.risque_service import creer_risque, dernieres_cotations, vers_sortie

COLONNES_REQUISES = {
    "N°",
    "Catégorie de risque",
    "Zone / Activité",
    "Danger identifié",
    "Personnes exposées",
    "P",
    "G",
    "Mesures de maîtrise proposées",
}
COLONNES_IGNOREES = {"C", "Niveau de risque", "Responsable", "Échéance", "Statut"}

_CATEGORIES_PAR_LIBELLE = {c.value.lower(): c for c in CategorieRisque}


def _lire_lignes_csv(contenu: bytes) -> list[dict]:
    texte = contenu.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(texte)))


def _lire_lignes_xlsx(contenu: bytes) -> list[dict]:
    from openpyxl import load_workbook

    classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    # Le classeur réel a plusieurs feuilles (Cotation, Registre, Synthèse) : on
    # cherche la première dont une ligne contient toutes les colonnes requises,
    # plutôt que de supposer le nom "Registre" ou la position de la feuille.
    for feuille in classeur.worksheets:
        lignes_brutes = list(feuille.iter_rows(values_only=True))
        for i, ligne in enumerate(lignes_brutes):
            cellules = {str(c).strip() for c in ligne if c is not None}
            if COLONNES_REQUISES <= cellules:
                entetes = [str(c).strip() if c is not None else "" for c in ligne]
                resultat = []
                for ligne_donnees in lignes_brutes[i + 1 :]:
                    if all(c is None for c in ligne_donnees):
                        continue
                    resultat.append(
                        {
                            entetes[j]: ("" if ligne_donnees[j] is None else str(ligne_donnees[j]).strip())
                            for j in range(len(entetes))
                            if entetes[j]
                        }
                    )
                return resultat
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=f"En-têtes attendus introuvables dans le fichier (colonnes requises : {sorted(COLONNES_REQUISES)})",
    )


def _construire_risque(ligne: dict) -> tuple[RisqueCreation | None, list[str]]:
    erreurs: list[str] = []

    manquants = [c for c in COLONNES_REQUISES if not (ligne.get(c) or "").strip()]
    if manquants:
        erreurs.append(f"Colonnes obligatoires manquantes : {', '.join(sorted(manquants))}")
        return None, erreurs

    categorie = _CATEGORIES_PAR_LIBELLE.get(ligne["Catégorie de risque"].strip().lower())
    if categorie is None:
        erreurs.append(
            f"Catégorie « {ligne['Catégorie de risque']} » inconnue "
            f"(valeurs acceptées : {', '.join(c.value for c in CategorieRisque)})"
        )

    try:
        probabilite = int(ligne["P"])
        if not (1 <= probabilite <= 5):
            raise ValueError
    except ValueError:
        erreurs.append(f"P « {ligne['P']} » doit être un entier entre 1 et 5")
        probabilite = None

    try:
        gravite = int(ligne["G"])
        if not (1 <= gravite <= 5):
            raise ValueError
    except ValueError:
        erreurs.append(f"G « {ligne['G']} » doit être un entier entre 1 et 5")
        gravite = None

    # Le classeur sépare "Danger identifié" (cause) et "Risque / Dommage
    # potentiel" (conséquence) ; le dictionnaire (7.2.2) ne modélise qu'un seul
    # champ `danger` (200 caractères) — concaténés plutôt que de perdre l'un
    # des deux silencieusement.
    danger_identifie = ligne["Danger identifié"].strip()
    dommage_potentiel = (ligne.get("Risque / Dommage potentiel") or "").strip()
    danger = f"{danger_identifie} — {dommage_potentiel}" if dommage_potentiel else danger_identifie
    if len(danger) > 200:
        erreurs.append(f"Danger + dommage potentiel dépasse 200 caractères ({len(danger)}) : à raccourcir manuellement")

    if erreurs:
        return None, erreurs

    return (
        RisqueCreation(
            danger=danger,
            categorie=categorie,
            unite_travail=ligne["Zone / Activité"].strip(),
            personnes_exposees=ligne["Personnes exposées"].strip() or None,
            cotation=CotationEntree(
                probabilite=probabilite,
                gravite=gravite,
                mesures_existantes=(ligne.get("Mesures de prévention existantes") or "").strip() or None,
                mesures_proposees=ligne["Mesures de maîtrise proposées"].strip(),
            ),
        ),
        [],
    )


def importer_registre(
    db: Session, contenu: bytes, nom_fichier: str, date_evaluation: date | None, auteur_id: int
) -> RapportImportRisques:
    """`date_evaluation` s'applique à TOUTES les lignes importées : le classeur
    réel ne donne qu'une période globale ("mars – mai 2026"), pas une date par
    ligne — par défaut, la date du jour de l'import."""
    if nom_fichier.lower().endswith(".xlsx"):
        lignes = _lire_lignes_xlsx(contenu)
    elif nom_fichier.lower().endswith(".csv"):
        lignes = _lire_lignes_csv(contenu)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Format non pris en charge : CSV ou XLSX attendu"
        )

    risques_crees: list[RisqueSortie] = []
    erreurs: list[LigneErreurImportRisque] = []

    for numero_ligne, ligne in enumerate(lignes, start=1):
        donnees, erreurs_ligne = _construire_risque(ligne)
        numero_source = (ligne.get("N°") or "").strip()
        numero_int = int(numero_source) if numero_source.isdigit() else None
        if donnees is None:
            erreurs.append(LigneErreurImportRisque(ligne=numero_ligne, numero=numero_int, erreurs=erreurs_ligne))
            continue
        if date_evaluation is not None:
            donnees.cotation.date_evaluation = date_evaluation
        risque = creer_risque(db, donnees, auteur_id=auteur_id)
        cotation = dernieres_cotations(db, [risque.id]).get(risque.id)
        risques_crees.append(vers_sortie(risque, cotation))

    return RapportImportRisques(
        total_lignes=len(lignes),
        importes=len(risques_crees),
        en_erreur=len(erreurs),
        risques=risques_crees,
        erreurs=erreurs,
    )
